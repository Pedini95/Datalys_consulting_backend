"""
Générateur de rapports pour les incidents
Formats supportés : PDF, Excel, CSV
"""

import csv
import io
import logging
import os
import base64
from datetime import datetime
from typing import List, Any
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

logger = logging.getLogger(__name__)

# Logo Datalys en base64
DATALYS_LOGO_BASE64 = "/9j/4AAQSkZJRgABAQAAAQABAAD/2wBDAAUDBAQEAwUEBAQFBQUGBwwIBwcHBw8LCwkMEQ8SEhEPERETFhwXExQaFRERGCEYGh0dHx8fExciJCIeJBweHx7/2wBDAQUFBQcGBw4ICA4eFBEUHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh4eHh7/wAARCABQAMgDASIAAhEBAxEB/8QAHAABAAICAwEAAAAAAAAAAAAAAAYHAgUBAwgE/8QAOxAAAQMDAgQDBQUGBwEAAAAAAQACAwQFEQYSByExQRMiUQgTYXGhFSMyUoEWM0KRsfA0Q3Sys8HR4f/EABsBAQACAwEBAAAAAAAAAAAAAAACBAMFBgEH/8QALxEAAgECAwUHAwUAAAAAAAAAAAECAxEEITEFEkFRYRMUcYGR0fAiocEjMlKx4f/aAAwDAQACEQMRAD8A9loiIAiIgCIiAIiIAiIgCIiAIiIAiIgCIiAIiIAiIgCIiAIiIAiErgnCrTiHxBjp5pLZZJt1QQWzVLDyjOPwtP5vj2+anTpyqO8StisVTw0N+o/8ATaa14gUVlr222kYKusaczNa7DYwBnaT+Y+nZfHQ8VrVIQKu3VcHLJc0teB/RVTZbPVXe7RU1A2aWoe7cd3mA583OPYepXVfqWWhuVTa5Hsa6mldG8HluIPX5dwtjHC0v28TlZ7Yxl3VjlG9llkemLfVQ11FDWU7t8M0bZI3erSMhd6gPBW7eJ0iKKeRoko5nRMJcPMw+YY9cZI/RT4HIytbNKM3G+h1mGrdvRjU5oIui4VHhaCoqdu/3MTpNucZwCcfRefeGvFjjbxD0nBqjTWgNJut08kkcfiLvIx+WO2nI2+qiZz0Sio248XeJGjB4viNwmqYbOznNc7FWitZC3u58eA4AevJWzovVFi1jp2mv+m7lDcLdUtzHLGeh7tcOrXDuDzCA3KImQgCKO8QNa6c0HY2XrU9f4KhfUx0zZBG5/wB48naMNBPYnPwUiBBGQeSAIi+W71TqK11VYxge6CF8gaTgEtaTj6ID6kVe+zvrq48R+E1p1ddaSmpaysdO2SOnz7sbJXMBGSTzDR3VhIAiIgCIiAIiIAuHkNaS4gAdSVytNriimuGlLjSU73NlfCS3acbiOe35HGP1XqV3YhUk4wckr2K+4h8QhUzyWOxSExuBZNVMdjcfysPp23fyUG0xYKm+XWOlt0MgkBDnybssjGfxOz/ZWua2BgEzowwNILTg564A9VYWk9f2qx0TaSOwuiidh7pIpg5zye7sgf/ABbZxdKFqSOHVeONr7+KnZfMl7lj6T03QadoTBSM3Sv5yzOHmkP/AEB2ChHETQtXdtXNulPE51JLC01AjI3l7eWB8xjn8FO9M3+kv1GamliniAxlszdrsHofktnPLHFE6WVwaxoySey006snGThOzzV1bLn0yOx7ph6tOEHG8VZpfOf3KxoaGSiAphRSQRtbhrfdkAYW1t9dPSVURE0hjDxvbu5Ed+SX/UM1bL7qn3RUrT0PIv8An8Pgs7Pb5rjNhpLYWnzvI+g+K+Ybr75u4ObnK+vN+3U7J2VD9WKiracvnIll+INiriDy8NJ/sKpX2D/8Olp/1lX/AMpV03WB32FVU8LXPd4Z7GNHMk7CAPmvLPsx8TIOH/CSh0rqLRGvDcaepqHyeGsMsjMPkLhz5dl9QV7ZnLnrORjHxuY9oc1wwQRkEehXmThFVW3hz7RnFrTlBUCDSNHbmXyWFn7ujk2xukDR0HKRwwOzWjspXcOLuv8AVEElBw14UagZUyeVty1DEKKlgz/GWk7n49B9V9/DTgjTWLRGqaHU11dedRawilF9ue3BcZGuG2MHo1pcSPU9gMAegjuiZ+LfGi2ftdHq5/D/AEtVSPNqoqClZNWTxA7RJLI/k3JB5D+mCY7x9uXG7hBw+ra6k1sNTWmpcyAXGejjirrZIXgh3lBbIxwBZkjILgtzwy1vqDgzpyn0BxJ0nfJaa1bobdfbRROq6Wqpw4lm4M8zHAHGCOgGfjF/af15qfiXwmutv0VorUEWnYHQzXK5XCjdA6cCRuyKCIje/wA5a5zsYAagMfbDt2rKzhNa9TVmtJpbVW1VtLbP4CIMimdDzkEo8x8072Dy82OwVuGr1zpA02mp9WHWOob3LmilqqCKmjoomt87nCP8Q78/QqI+1PZbzdPZw07b7baa6trI6y2F9PBA58jQIyCS0DIwSAfRWZxFsd4bfrLrLT9MKyvtYdFNRl2DPC4eYNP5hk/zWfDKLqLe666XtlfzK+Kc1Se5001tfO3kZDSetNniHcQqzxuM7RRx+4z6bOuPqozZtRarrNXansGonxR+CscmYYW/dueAPvW98ODs4UmHE22GH3ZsWoxW4x4X7Ofv3emen65UR0/FqKu4k6quF3tUtLPV2B3uoGjdsaQBHHuHIvwOYHclXFGbpz7WK0yySeq0sUXKmqlNUZN555trR63Ki4Uas1hoD2ZeH2trNtqtM2+41cepKBsAdI+nkqXtEzXdRsPYd3DPLK9V3LVlhoNEy6xmuERskVF441LTlrodu4Ob6kjGB3JAVXeyXpyVvsy2rT2p7RNC2obWRVVHWQljnRyTyZDmnmAWn6qntNaX1Pedb1Ps2SV7a7Q+nLqLnW1schMhoTiSKieex3u5jqDk9GrWG2LDpdd8XZ/Z01TxLipGvuNwmNXp+3Cla59DQF4aHkDnIdmX889AehwsOD8uu7rcNMaj0txip9c2mqLf2jt9wbHG+lYWguMbWjcxzTkbTjnjqCrT4tah1DobSNJc9J6R+34KSoijrKKmJEsVIBhzomNHmLcABvpz6BefNcV2n9Y8RNKXjgfpm927WTLmx9yrY7XJRUzKb/MbU7gGuPTPXIyOeQgPXyIOiIAiIgCEZREBX944V2OtmfLT1dZSOcSQ0EOY3JzgAjkFoajh4yySRTVFa2tiDiIw6LaAevm5nPfCt5fBfKM11vkhZj3nJzMnuFg2lUxNTCThSk07Zc/Dz0K2H2bgo4iNSVNZPy9NCIaVqDbrk+SZw9y+MhxHPmOY5f31XVfbvWXCoa4xyRUzHZbGO/xPqVhW2u8QvD5KKUtacjZ5h9FtbFZZKtwmqA+OAHk05Dn/APgXzihDaFWmsBGLSvfivW/A7CpLDwl3htN2Oiz2ma4SbnEspwebyPxfAKZ0lPDTQNhgYGMaMABZxRsjjaxjQ1rRgAdAsl3Oy9lUsBD6c5PV/OBosTip4iWenIJhEW1KoREQBMIiA4JA6lYRyxyAmN7HgHB2nK199/HTOma91GHn34aCe3lJx/DnqumpqqZlIPs58bIzK0TyQt/dtP8AF9AM9lSqYtQnJPh6vw6dfEyxpOSVuJuf1QAdloJ6qrZBW+EqZJoWNYY5SMkOLsOaDjzDH8srumfV01XPC2tfsNKZBJM3IY/djsOnwUe/x/i/t1XPoS7B8/mXuYcQLherVoq73DTlqfdrvBSvdRUjSMyy4w0c+wJyR3AIUL9nXQM3D7Qjpr/K2bU97nNwvlU92XPqJDnZnuG7semS4jqpgyvnbRmTfL93URiV5cJGbCeeDgcvX0SuqjVNrGscZYGupyzDcjO/nj16KEtpQ3LxWdvw3+CSw0t6z+ae5vy4AEkgAcyfRGOa5oc1wc0jIIOQVo3kxXO4tdVS73R7o43NBDxsPTlzwVlSiqqJWQeJmgjFHE/EbA3DjnPb6KUcc3Ld3c7tcODI9hZXubxcZUd+0KkxURqZ5GCWnLnbBscXg4ycjpjstpY8/Z8bjUCoe7JfIHZBd3UqGOjXnuRXC/8AXueVKLgrs+9ERXjCEREAREQBERAEREAREQBERAEREAXAAXKIAAAMBMIiWAAAGEAwiJYDHPKIiA+OpohJUiojlkikLdji08nDtkfDK7aWmjpotkeeZLnEnJcT1JXeixRowjJzSzZJybVmERFlIn//2Q=="


class ReportGenerator:
    """Classe pour générer des rapports d'incidents dans différents formats"""

    def __init__(self):
        self.styles = getSampleStyleSheet()

    def generate_csv(self, incidents: List[Any], include_stats: bool = False) -> str:
        """
        Génère un rapport CSV des incidents

        Args:
            incidents: Liste des incidents
            include_stats: Inclure les statistiques

        Returns:
            str: Contenu CSV
        """
        try:
            output = io.StringIO()
            writer = csv.writer(output)

            # En-têtes
            headers = [
                'Numéro',
                'Titre',
                'Description',
                'Statut',
                'Priorité',
                'Impact',
                'Domaine',
                'Déclarant',
                'Assigné à',
                'Créé le',
                'Résolu le',
                'Notes de résolution',
                'Nombre de refus'
            ]
            writer.writerow(headers)

            # Données
            for incident in incidents:
                row = [
                    incident.incident_number or '',
                    incident.title or '',
                    (incident.description or '')[:100] + '...' if incident.description and len(incident.description) > 100 else (incident.description or ''),
                    incident.status or '',
                    incident.priority or '',
                    incident.impact or '',
                    incident.domain or '',
                    incident.declarant_name or '',
                    str(incident.assigned_to) if incident.assigned_to else '',
                    incident.created_at.strftime('%Y-%m-%d %H:%M') if incident.created_at else '',
                    incident.resolved_at.strftime('%Y-%m-%d %H:%M') if incident.resolved_at else '',
                    (incident.resolution_notes or '')[:100] + '...' if incident.resolution_notes and len(incident.resolution_notes) > 100 else (incident.resolution_notes or ''),
                    str(incident.refusal_count or 0)
                ]
                writer.writerow(row)

            # Statistiques si demandées
            if include_stats:
                writer.writerow([])
                writer.writerow(['=== STATISTIQUES ==='])
                writer.writerow(['Total incidents', len(incidents)])

                # Comptage par statut
                stats_by_status = {}
                stats_by_priority = {}
                stats_by_domain = {}

                for incident in incidents:
                    status = incident.status or 'inconnu'
                    priority = incident.priority or 'inconnu'
                    domain = incident.domain or 'non_specifie'

                    stats_by_status[status] = stats_by_status.get(status, 0) + 1
                    stats_by_priority[priority] = stats_by_priority.get(priority, 0) + 1
                    stats_by_domain[domain] = stats_by_domain.get(domain, 0) + 1

                writer.writerow([])
                writer.writerow(['Par statut'])
                for status, count in stats_by_status.items():
                    writer.writerow([status, count])

                writer.writerow([])
                writer.writerow(['Par priorité'])
                for priority, count in stats_by_priority.items():
                    writer.writerow([priority, count])

                writer.writerow([])
                writer.writerow(['Par domaine'])
                for domain, count in stats_by_domain.items():
                    writer.writerow([domain, count])

            return output.getvalue()

        except Exception as e:
            logger.error(f"Erreur génération CSV : {str(e)}")
            raise

    def generate_excel(self, incidents: List[Any], include_stats: bool = False) -> bytes:
        """
        Génère un rapport Excel des incidents

        Args:
            incidents: Liste des incidents
            include_stats: Inclure les statistiques

        Returns:
            bytes: Contenu Excel
        """
        try:
            # Import openpyxl uniquement pour la génération Excel
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Incidents"

            # Styles
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # En-têtes
            headers = [
                'Numéro', 'Titre', 'Description', 'Statut', 'Priorité',
                'Impact', 'Domaine', 'Déclarant', 'Assigné à',
                'Créé le', 'Résolu le', 'Notes de résolution', 'Refus'
            ]

            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border

            # Données
            for row_num, incident in enumerate(incidents, 2):
                data = [
                    incident.incident_number or '',
                    incident.title or '',
                    (incident.description or '')[:200] + '...' if incident.description and len(incident.description) > 200 else (incident.description or ''),
                    incident.status or '',
                    incident.priority or '',
                    incident.impact or '',
                    incident.domain or '',
                    incident.declarant_name or '',
                    str(incident.assigned_to) if incident.assigned_to else '',
                    incident.created_at.strftime('%Y-%m-%d %H:%M') if incident.created_at else '',
                    incident.resolved_at.strftime('%Y-%m-%d %H:%M') if incident.resolved_at else '',
                    (incident.resolution_notes or '')[:200] + '...' if incident.resolution_notes and len(incident.resolution_notes) > 200 else (incident.resolution_notes or ''),
                    incident.refusal_count or 0
                ]

                for col_num, value in enumerate(data, 1):
                    cell = ws.cell(row=row_num, column=col_num, value=value)
                    cell.border = border

                    # Coloration par statut
                    if col_num == 4:  # Colonne Statut
                        if incident.status == 'resolu':
                            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                        elif incident.status == 'en_cours':
                            cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                        elif incident.status == 'nouveau':
                            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            # Ajuster la largeur des colonnes
            for col_num in range(1, len(headers) + 1):
                column_letter = get_column_letter(col_num)
                ws.column_dimensions[column_letter].width = 20

            # Ajouter les statistiques si demandées
            if include_stats:
                self._add_excel_stats(wb, incidents)

            # Sauvegarder dans un buffer
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            return output.getvalue()

        except Exception as e:
            logger.error(f"Erreur génération Excel : {str(e)}")
            raise

    def _add_excel_stats(self, wb, incidents: List[Any]):
        """Ajoute une feuille de statistiques dans le workbook Excel"""
        # Import openpyxl pour les styles
        from openpyxl.styles import Font

        ws_stats = wb.create_sheet("Statistiques")

        # Titre
        ws_stats['A1'] = 'STATISTIQUES DES INCIDENTS'
        ws_stats['A1'].font = Font(size=16, bold=True, color="4472C4")

        # Total
        ws_stats['A3'] = 'Total incidents'
        ws_stats['B3'] = len(incidents)
        ws_stats['B3'].font = Font(bold=True)

        # Statistiques par statut
        stats_by_status = {}
        stats_by_priority = {}
        stats_by_domain = {}

        for incident in incidents:
            status = incident.status or 'inconnu'
            priority = incident.priority or 'inconnu'
            domain = incident.domain or 'non_specifie'

            stats_by_status[status] = stats_by_status.get(status, 0) + 1
            stats_by_priority[priority] = stats_by_priority.get(priority, 0) + 1
            stats_by_domain[domain] = stats_by_domain.get(domain, 0) + 1

        # Par statut
        row = 5
        ws_stats[f'A{row}'] = 'Par statut'
        ws_stats[f'A{row}'].font = Font(bold=True)
        row += 1
        for status, count in stats_by_status.items():
            ws_stats[f'A{row}'] = status
            ws_stats[f'B{row}'] = count
            row += 1

        # Par priorité
        row += 1
        ws_stats[f'A{row}'] = 'Par priorité'
        ws_stats[f'A{row}'].font = Font(bold=True)
        row += 1
        for priority, count in stats_by_priority.items():
            ws_stats[f'A{row}'] = priority
            ws_stats[f'B{row}'] = count
            row += 1

        # Par domaine
        row += 1
        ws_stats[f'A{row}'] = 'Par domaine'
        ws_stats[f'A{row}'].font = Font(bold=True)
        row += 1
        for domain, count in stats_by_domain.items():
            ws_stats[f'A{row}'] = domain
            ws_stats[f'B{row}'] = count
            row += 1

        # Ajuster largeur
        ws_stats.column_dimensions['A'].width = 25
        ws_stats.column_dimensions['B'].width = 15

    def generate_pdf(self, incidents: List[Any], include_stats: bool = False) -> bytes:
        """
        Génère un rapport PDF des incidents

        Args:
            incidents: Liste des incidents
            include_stats: Inclure les statistiques

        Returns:
            bytes: Contenu PDF
        """
        try:
            buffer = io.BytesIO()
            doc = SimpleDocTemplate(
                buffer,
                pagesize=landscape(A4),
                rightMargin=30,
                leftMargin=30,
                topMargin=30,
                bottomMargin=18
            )

            elements = []

            # Logo Datalys (fichier physique en priorité, base64 en secours)
            try:
                # 1. Essayer de charger un fichier logo depuis src/static/image
                project_root = os.getcwd()
                logo_path = os.path.join(project_root, 'src', 'static', 'image', 'logo_datalys.png')

                if os.path.exists(logo_path):
                    logger.info(f"Chargement du logo PDF depuis le fichier: {logo_path}")
                    logo = Image(logo_path, width=2*inch, height=0.8*inch)
                else:
                    # 2. Fallback: utiliser le logo en base64 embarqué
                    logger.info("Logo fichier introuvable, utilisation du logo base64 pour le PDF")
                    logo_data = base64.b64decode(DATALYS_LOGO_BASE64)
                    logo_buffer = io.BytesIO(logo_data)
                    logo = Image(ImageReader(logo_buffer), width=2*inch, height=0.8*inch)

                elements.append(logo)
                elements.append(Spacer(1, 15))
            except Exception as e:
                logger.warning(f"Erreur lors du chargement du logo pour PDF: {str(e)}")

            # Titre
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=self.styles['Heading1'],
                fontSize=18,
                textColor=colors.HexColor('#4472C4'),
                spaceAfter=30,
                alignment=1  # Centre
            )
            elements.append(Paragraph(f"RAPPORT D'INCIDENTS", title_style))
            elements.append(Paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}", self.styles['Normal']))
            elements.append(Spacer(1, 20))

            # Statistiques en haut si demandées
            if include_stats:
                stats_data = self._generate_stats_data(incidents)
                stats_table = Table(stats_data, colWidths=[2*inch, 1.5*inch])
                stats_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                elements.append(stats_table)
                elements.append(Spacer(1, 20))

            # Tableau des incidents
            data = [['N°', 'Titre', 'Statut', 'Priorité', 'Domaine', 'Créé le', 'Résolu le']]

            for incident in incidents:
                row = [
                    incident.incident_number or '',
                    (incident.title or '')[:40] + '...' if incident.title and len(incident.title) > 40 else (incident.title or ''),
                    incident.status or '',
                    incident.priority or '',
                    incident.domain or '',
                    incident.created_at.strftime('%d/%m/%Y') if incident.created_at else '',
                    incident.resolved_at.strftime('%d/%m/%Y') if incident.resolved_at else '-'
                ]
                data.append(row)

            table = Table(data, colWidths=[1.2*inch, 2.5*inch, 1*inch, 0.8*inch, 1.2*inch, 1*inch, 1*inch])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey])
            ]))

            elements.append(table)

            # Construire le PDF
            doc.build(elements)
            buffer.seek(0)

            return buffer.getvalue()

        except Exception as e:
            logger.error(f"Erreur génération PDF : {str(e)}")
            raise

    def _generate_stats_data(self, incidents: List[Any]) -> List[List[str]]:
        """Génère les données statistiques pour le PDF"""
        stats_by_status = {}
        stats_by_priority = {}

        for incident in incidents:
            status = incident.status or 'inconnu'
            priority = incident.priority or 'inconnu'

            stats_by_status[status] = stats_by_status.get(status, 0) + 1
            stats_by_priority[priority] = stats_by_priority.get(priority, 0) + 1

        data = [['Statistique', 'Valeur']]
        data.append(['Total incidents', str(len(incidents))])
        data.append(['', ''])  # Ligne vide

        for status, count in stats_by_status.items():
            data.append([f'Statut: {status}', str(count)])

        data.append(['', ''])  # Ligne vide

        for priority, count in stats_by_priority.items():
            data.append([f'Priorité: {priority}', str(count)])

        return data
